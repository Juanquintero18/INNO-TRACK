"""Motor de importación masiva desde Excel o CSV.

Cada módulo describe su importación con un `ImportSpec` (columnas, alias de
encabezados, ejemplo de plantilla, validación de una fila y guardado). El
motor se encarga de lo común:

- leer archivos .csv, .xls, .xlsx y .xlsm;
- reconocer encabezados escritos de forma distinta (acentos, espacios, alias
  o errores de tipeo);
- construir una vista previa fila por fila sin guardar nada;
- confirmar el lote completo en una sola transacción;
- generar la plantilla de ejemplo.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from difflib import get_close_matches
from io import BytesIO, StringIO
import re
import unicodedata
from typing import Any

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
import xlrd


DATE_FORMATS = (
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%Y/%m/%d",
)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SourceRow = tuple[int, dict[str, object]]


# --- Normalización de texto ------------------------------------------------------


def normalize_text(value: object) -> str:
    """Normaliza texto para comparaciones tolerantes a mayúsculas/acentos."""

    if value is None:
        return ""
    text = str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).casefold()


def normalize_lookup_key(value: object) -> str:
    """Genera una clave de búsqueda estable removiendo ruido común."""

    normalized = normalize_text(value)
    normalized = re.sub(r"[-_/\\.]+", " ", normalized)
    normalized = re.sub(r"[^a-z0-9\s]", "", normalized)
    return re.sub(r"\s+", " ", normalized).strip()


def compact_lookup_key(value: object) -> str:
    """Versión compacta de la clave para matching flexible."""

    return normalize_lookup_key(value).replace(" ", "")


def cell_to_text(value: object) -> str:
    """Convierte cualquier celda a representación textual consistente."""

    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


# --- Conversión de valores -------------------------------------------------------


def parse_decimal(value: object, label: str = "La cantidad", masculine: bool = False) -> Decimal:
    """Parsea números permitiendo distintos formatos (coma o punto decimal).

    `masculine` concuerda los mensajes con etiquetas como "El costo" en vez de "La cantidad".
    """
    obligatorio, numerico = ("obligatorio", "numérico") if masculine else ("obligatoria", "numérica")

    if value is None or str(value).strip() == "":
        raise ValueError(f"{label} es {obligatorio}.")

    if isinstance(value, Decimal):
        return value

    if isinstance(value, (int, float)):
        return Decimal(str(value))

    text = str(value).strip().replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")

    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{label} debe ser {numerico}.") from exc


def parse_date(value: object, label: str = "La fecha") -> str:
    """Parsea fecha y retorna formato ISO para persistencia."""

    if value is None or str(value).strip() == "":
        raise ValueError(f"{label} es obligatoria.")

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    text = str(value).strip()
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date().isoformat()
        except ValueError:
            continue

    raise ValueError(f"{label} debe usar formato YYYY-MM-DD o DD/MM/YYYY.")


def parse_optional_date(value: object, label: str) -> str | None:
    """Como `parse_date`, pero una celda vacía es válida y devuelve `None`."""

    if value is None or str(value).strip() == "":
        return None
    return parse_date(value, label)


def email_format_error(value: str) -> str | None:
    """Mensaje de error si el texto no parece un correo electrónico."""
    if "@" not in value or "." not in value.split("@")[-1] or " " in value:
        return f"El correo '{value}' no es válido."
    return None


def flatten_serializer_errors(detail: object) -> list[str]:
    """Aplana errores anidados de serializer a una lista simple de mensajes."""

    if isinstance(detail, dict):
        messages: list[str] = []
        for value in detail.values():
            messages.extend(flatten_serializer_errors(value))
        return messages
    if isinstance(detail, list):
        messages: list[str] = []
        for item in detail:
            messages.extend(flatten_serializer_errors(item))
        return messages
    return [str(detail)]


# --- Catálogos de referencia -----------------------------------------------------


class CatalogLookup:
    """Catálogo en memoria para resolver textos del archivo a registros del sistema.

    Cada registro se indexa bajo varias normalizaciones de su nombre (sin
    acentos, sin signos, sin espacios) para tolerar diferencias de escritura.
    """

    def __init__(self, not_found_message: str):
        self.not_found_message = not_found_message
        self._instances: dict[str, object] = {}
        self._display_names: dict[str, str] = {}

    @staticmethod
    def _keys(raw_value: object) -> tuple[str, str, str]:
        return normalize_text(raw_value), normalize_lookup_key(raw_value), compact_lookup_key(raw_value)

    def register(self, raw_value: object, instance: object) -> None:
        """Registra múltiples variantes de un valor para resolución flexible."""
        display_value = cell_to_text(raw_value)

        for key in set(self._keys(raw_value)):
            if not key:
                continue
            # Ante nombres equivalentes gana el primer registro cargado.
            self._instances.setdefault(key, instance)
            self._display_names.setdefault(key, display_value)

    def find(self, raw_value: object):
        """Busca una referencia de catálogo usando distintas normalizaciones."""
        for key in self._keys(raw_value):
            if key and key in self._instances:
                return self._instances[key]
        return None

    def suggest(self, raw_value: object) -> str | None:
        """Sugiere el nombre más cercano para mejorar mensajes de validación."""
        compact_target = compact_lookup_key(raw_value)
        if not compact_target:
            return None

        compact_candidates = [key for key in self._display_names if key and " " not in key]
        matches = get_close_matches(compact_target, compact_candidates, n=1, cutoff=0.8)
        if matches:
            return self._display_names[matches[0]]

        normalized_candidates = [key for key in self._display_names if key and " " in key]
        matches = get_close_matches(normalize_lookup_key(raw_value), normalized_candidates, n=1, cutoff=0.8)
        if matches:
            return self._display_names[matches[0]]

        return None

    def not_found_error(self, raw_value: object) -> str:
        """Mensaje de error para un valor ausente del catálogo, con sugerencia si la hay."""
        suggestion = self.suggest(raw_value)

        if suggestion:
            return f"{self.not_found_message} ¿Quisiste decir '{suggestion}'?"
        return self.not_found_message


class DuplicateTracker:
    """Detecta valores que ya existen en el sistema o que se repiten dentro del archivo."""

    def __init__(self, existing_values, label: str):
        self.label = label
        self._existing = {normalize_text(value) for value in existing_values if value}
        self._seen_in_file: set[str] = set()

    def check(self, value: str) -> str | None:
        """Mensaje de error si el valor está repetido; `None` si es nuevo. Lo registra como visto."""
        key = normalize_text(value)
        if not key:
            return None
        if key in self._existing:
            return f"Ya existe un registro con {self.label} '{value}'."
        if key in self._seen_in_file:
            return f"{self.label.capitalize()} '{value}' está repetido en el archivo."
        self._seen_in_file.add(key)
        return None


# --- Especificación de una importación ------------------------------------------


@dataclass
class RowResult:
    """Resultado de validar una fila: errores, datos a guardar y nombres reconocidos."""

    errors: list[str] = field(default_factory=list)
    payload: dict[str, Any] | None = None
    resolved: dict[str, str | None] = field(default_factory=dict)


class ImportSpec:
    """Describe la importación de un tipo de registro. Cada módulo define la suya."""

    #: Columnas canónicas, en el orden de la plantilla y de la vista previa.
    headers: list[str] = []
    #: Variantes de encabezado aceptadas -> columna canónica.
    aliases: dict[str, str] = {}
    sheet_title: str = "datos"
    template_file_name: str = "plantilla.xlsx"
    example_rows: list[list[object]] = []

    def new_context(self) -> Any:
        """Estado compartido durante el análisis de un archivo (catálogos, repetidos...)."""
        return None

    def validate_row(self, values: dict[str, str], raw_row: dict[str, object], context: Any) -> RowResult:
        """Valida una fila. `values` trae el texto de cada columna y `raw_row` las celdas originales."""
        raise NotImplementedError

    def save(self, payload: dict[str, Any], user) -> Any:
        """Crea el registro de una fila válida y lo devuelve."""
        raise NotImplementedError


# --- Lectura de archivos -----------------------------------------------------------


def _build_header_observation(raw_text: str, canonical: str) -> str | None:
    """Informa cuándo una cabecera fue reinterpretada por alias/fuzzy."""

    if not raw_text:
        return None
    if normalize_lookup_key(raw_text) == normalize_lookup_key(canonical):
        return None
    return f"Se interpretó '{raw_text}' como '{canonical}'."


def resolve_header(value: object, spec: ImportSpec) -> tuple[str, str | None]:
    """Resuelve cabecera de archivo hacia la cabecera canónica del sistema."""

    canonical_headers = set(spec.headers)
    aliases = spec.aliases
    raw_text = cell_to_text(value)
    # Estrategia en capas: exacto -> alias -> fuzzy para soportar plantillas ruidosas.
    candidates = {
        normalize_text(raw_text),
        normalize_lookup_key(raw_text),
        compact_lookup_key(raw_text),
    }

    for candidate in list(candidates):
        if not candidate:
            continue

        variants = {
            candidate,
            candidate.replace("-", "_"),
            candidate.replace(" ", "_"),
            candidate.replace("_", " "),
            candidate.replace("_", ""),
        }

        for variant in variants:
            if variant in canonical_headers:
                return variant, _build_header_observation(raw_text, variant)
            if variant in aliases:
                canonical = aliases[variant]
                return canonical, _build_header_observation(raw_text, canonical)

    alias_candidates = set(aliases) | canonical_headers
    normalized_target = normalize_lookup_key(raw_text)
    compact_target = compact_lookup_key(raw_text)

    normalized_matches = get_close_matches(normalized_target, alias_candidates, n=1, cutoff=0.72)
    if normalized_matches:
        match = normalized_matches[0]
        canonical = aliases.get(match, match)
        return canonical, _build_header_observation(raw_text, canonical)

    compact_aliases = {compact_lookup_key(alias): aliases.get(alias, alias) for alias in alias_candidates}
    compact_matches = get_close_matches(compact_target, list(compact_aliases), n=1, cutoff=0.8)
    if compact_matches:
        canonical = compact_aliases[compact_matches[0]]
        return canonical, _build_header_observation(raw_text, canonical)

    fallback = normalize_text(raw_text).replace("-", "_").replace(" ", "_")
    canonical = aliases.get(fallback, fallback)
    return canonical, _build_header_observation(raw_text, canonical)


def _rows_from_matrix(rows: list[list[object]], spec: ImportSpec) -> tuple[list[SourceRow], list[str]]:
    """Mapea matriz cruda a filas numeradas con cabeceras normalizadas."""

    # Ignora filas vacías para estabilizar la numeración de errores por archivo.
    non_empty_rows = [row for row in rows if any(cell_to_text(cell) for cell in row)]
    if not non_empty_rows:
        return [], []

    header_row = non_empty_rows[0]
    resolved_headers = [resolve_header(cell, spec) for cell in header_row]
    headers = [canonical for canonical, _observation in resolved_headers]
    header_observations = list(dict.fromkeys(observation for _canonical, observation in resolved_headers if observation))
    parsed_rows: list[SourceRow] = []

    for row_offset, row in enumerate(non_empty_rows[1:], start=2):
        # Se rellena con None cuando faltan columnas para mantener mapeo por cabecera.
        values = row + [None] * max(0, len(headers) - len(row))
        parsed_rows.append(
            (
                row_offset,
                {header: values[index] if index < len(values) else None for index, header in enumerate(headers) if header},
            )
        )

    return parsed_rows, header_observations


def _read_matrix(file_name: str, file_bytes: bytes) -> list[list[object]]:
    """Lee el archivo como una matriz de celdas según su extensión."""

    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if extension == "csv":
        try:
            decoded_text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded_text = file_bytes.decode("latin-1")
        return list(csv.reader(StringIO(decoded_text)))

    if extension in {"xlsx", "xlsm"}:
        worksheet = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True).active
        return [list(row) for row in worksheet.iter_rows(values_only=True)]

    if extension == "xls":
        workbook = xlrd.open_workbook(file_contents=file_bytes)
        worksheet = workbook.sheet_by_index(0)
        rows: list[list[object]] = []

        for row_index in range(worksheet.nrows):
            row_values: list[object] = []
            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                value: object = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode)
                row_values.append(value)
            rows.append(row_values)

        return rows

    raise ValueError("El archivo debe estar en formato .xlsx, .xls o .csv.")


# --- Análisis, confirmación y plantilla -------------------------------------------


def analyze_import(spec: ImportSpec, file_name: str, file_bytes: bytes) -> tuple[dict[str, object], list[dict[str, Any]]]:
    """Analiza el archivo completo y produce la vista previa y los datos válidos.

    La vista previa y la confirmación usan exactamente este análisis, para que
    no haya diferencias entre lo que se muestra y lo que se guarda.
    """

    source_rows, header_observations = _rows_from_matrix(_read_matrix(file_name, file_bytes), spec)
    context = spec.new_context()
    preview_rows: list[dict[str, object]] = []
    payloads: list[dict[str, Any]] = []
    file_errors: list[str] = []

    if not source_rows:
        file_errors.append("El archivo no contiene filas de datos para importar.")

    for row_number, raw_row in source_rows:
        # La vista previa conserva valores originales del archivo para explicar errores al usuario.
        values = {header: cell_to_text(raw_row.get(header)) for header in spec.headers}
        result = spec.validate_row(values, raw_row, context)
        is_valid = not result.errors and result.payload is not None

        preview_rows.append(
            {
                "row_number": row_number,
                "status": "valid" if is_valid else "error",
                "errors": result.errors,
                "values": values,
                "resolved": result.resolved,
            }
        )
        if is_valid:
            payloads.append(result.payload)

    invalid_rows = sum(1 for row in preview_rows if row["status"] == "error")
    preview = {
        "file_name": file_name,
        "file_errors": file_errors,
        "header_observations": header_observations,
        "total_rows": len(preview_rows),
        "valid_rows": len(payloads),
        "invalid_rows": invalid_rows,
        "can_import": not file_errors and invalid_rows == 0 and bool(payloads),
        "rows": preview_rows,
    }
    return preview, payloads


@transaction.atomic
def run_import(spec: ImportSpec, file_name: str, file_bytes: bytes, user) -> dict[str, object]:
    """Importa el lote completo con rollback ante cualquier error."""

    preview, payloads = analyze_import(spec, file_name, file_bytes)

    if not preview["can_import"]:
        return {"created_count": 0, "preview": preview}

    # Si una fila falla al guardar, la transacción revierte todo el lote.
    created_ids = [spec.save(payload, user).id for payload in payloads]

    return {"created_count": len(created_ids), "created_ids": created_ids, "preview": preview}


def build_template_response(spec: ImportSpec) -> HttpResponse:
    """Genera la plantilla Excel de referencia con los encabezados y filas de ejemplo."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = spec.sheet_title
    worksheet.append(spec.headers)
    for row in spec.example_rows:
        worksheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)

    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{spec.template_file_name}"'
    return response
