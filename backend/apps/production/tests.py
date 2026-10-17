"""Pruebas de caracterización de la API de producción."""

from io import BytesIO

from openpyxl import load_workbook

from apps.production.models import Pieza, PiezaHistorial, PiezaMateriaPrima
from config.testing import ApiTestCase, csv_file


BASE = "/api/production"


class ProyectosYOrdenesTests(ApiTestCase):
    def test_listados(self):
        self.assert_golden("proyectos", self.client.get(f"{BASE}/proyectos/"))
        self.assert_golden("ordenes", self.client.get(f"{BASE}/ordenes/"))

    def test_solo_el_administrador_escribe(self):
        resultado = {}

        for rol, usuario in (("almacenista", self.almacenista), ("trabajador", self.trabajador)):
            self.login_as(usuario)
            resultado[f"{rol} crea proyecto"] = self.client.post(f"{BASE}/proyectos/", {"nombre": "X"}, format="json").status_code
            resultado[f"{rol} crea orden"] = self.client.post(f"{BASE}/ordenes/", {"codigo_orden": "X"}, format="json").status_code
            resultado[f"{rol} lista proyectos"] = self.client.get(f"{BASE}/proyectos/").status_code

        self.assert_golden_data("matriz", resultado)

    def test_crud_de_proyecto(self):
        creado = self.client.post(f"{BASE}/proyectos/", {"nombre": "Río Lento", "codigo": "RL-05", "descripcion": None, "fecha_inicio": "2026-04-01", "fecha_fin": None, "estado": "pausado"}, format="json")
        self.assert_golden("crear", creado)

        url = f"{BASE}/proyectos/{creado.json()['id']}/"
        self.assert_golden("actualizar", self.client.put(url, {"nombre": "Río Lento II", "codigo": None, "estado": None}, format="json"))
        self.assert_golden("eliminar", self.client.delete(url))

    def test_no_elimina_proyecto_con_ordenes(self):
        self.assert_golden("bloqueado", self.client.delete(f"{BASE}/proyectos/{self.aquaventura.id}/"))

    def test_crud_de_orden(self):
        creada = self.client.post(f"{BASE}/ordenes/", {"proyecto_id": self.sin_ordenes.id, "codigo_orden": "ORD-300", "fecha_creacion": "2026-03-14", "estado": "abierta"}, format="json")
        self.assert_golden("crear", creada)

        url = f"{BASE}/ordenes/{creada.json()['id']}/"
        self.assert_golden("mover_de_proyecto", self.client.put(url, {"proyecto_id": self.aquaventura.id, "codigo_orden": "ORD-300", "fecha_creacion": None, "estado": "cerrada"}, format="json"))
        self.assert_golden("eliminar", self.client.delete(url))

    def test_no_elimina_orden_con_piezas(self):
        self.assert_golden("bloqueada", self.client.delete(f"{BASE}/ordenes/{self.orden_1.id}/"))


class PiezasTests(ApiTestCase):
    def nueva_pieza(self, **cambios):
        return {
            "orden_id": self.orden_vacia.id,
            "usuario_id": self.trabajador.id,
            "trace_id": "TR-0050",
            "nombre": "Curva doble",
            "fecha_gelcoat": "2026-03-12",
            "fecha_qc": None,
            "peso_real": 30.25,
            "materias_primas": [
                {"materia_prima_id": self.resina.id, "cantidad_teorica": 12.5, "cantidad_real": 12.5},
                {"materia_prima_id": self.gelcoat.id, "cantidad_teorica": 2, "cantidad_real": None},
            ],
            **cambios,
        }

    def test_listado_con_historial_antiguo(self):
        self.assert_golden("listado", self.client.get(f"{BASE}/piezas/"))

    def test_crear_registra_historial_de_creacion(self):
        self.login_as(self.trabajador)

        response = self.client.post(f"{BASE}/piezas/", self.nueva_pieza(), format="json")

        self.assert_golden("crear", response)
        evento = PiezaHistorial.objects.get(pieza_id=response.json()["id"])
        self.assert_golden_data("historial_guardado", {"accion": evento.accion, "usuario_id": evento.usuario_id, "descripcion": evento.descripcion})

    def test_crear_sin_materiales(self):
        self.assert_golden("crear", self.client.post(f"{BASE}/piezas/", self.nueva_pieza(materias_primas=[]), format="json"))

    def test_editar_registra_cambios_de_campos_y_materiales(self):
        creada = self.client.post(f"{BASE}/piezas/", self.nueva_pieza(), format="json").json()
        url = f"{BASE}/piezas/{creada['id']}/"

        cambios = self.nueva_pieza(
            nombre="Curva doble reforzada",
            orden_id=self.orden_1.id,
            usuario_id=self.admin.id,
            fecha_qc="2026-03-20",
            peso_real=31,
            materias_primas=[
                {"materia_prima_id": self.resina.id, "cantidad_teorica": 12.5, "cantidad_real": 13},
                {"materia_prima_id": self.fibra.id, "cantidad_teorica": 6, "cantidad_real": 6},
            ],
        )
        self.assert_golden("editar_con_cambios", self.client.put(url, cambios, format="json"))
        self.assert_golden("editar_sin_cambios", self.client.put(url, cambios, format="json"))

        solo_campos = {"nombre": "Curva doble final"}
        self.assert_golden("editar_solo_un_campo_conserva_materiales", self.client.patch(url, solo_campos, format="json"))

    def test_materiales_deshabilitados(self):
        nueva = self.nueva_pieza(materias_primas=[{"materia_prima_id": self.catalizador.id, "cantidad_teorica": 1, "cantidad_real": 1}])
        self.assert_golden("no_se_agregan_a_piezas_nuevas", self.client.post(f"{BASE}/piezas/", nueva, format="json"))

        # La pieza de ejemplo ya tenía el catalizador: puede conservarlo al editarse.
        conservar = {
            "materias_primas": [
                {"materia_prima_id": self.resina.id, "cantidad_teorica": 20, "cantidad_real": 21.5},
                {"materia_prima_id": self.catalizador.id, "cantidad_teorica": 1, "cantidad_real": 0.8},
            ]
        }
        self.assert_golden("se_conservan_en_piezas_que_ya_los_tenian", self.client.patch(f"{BASE}/piezas/{self.pieza.id}/", conservar, format="json"))

    def test_permisos(self):
        self.login_as(self.almacenista)
        self.assertEqual(self.client.post(f"{BASE}/piezas/", self.nueva_pieza(), format="json").status_code, 403)
        self.assertEqual(self.client.get(f"{BASE}/piezas/").status_code, 200)

    def test_eliminar_borra_hijos_y_deja_auditoria_completa(self):
        self.assert_golden("eliminar", self.client.delete(f"{BASE}/piezas/{self.pieza.id}/"))

        self.assertFalse(Pieza.objects.filter(pk=self.pieza.id).exists())
        self.assertEqual(PiezaMateriaPrima.objects.count(), 0)
        self.assertEqual(PiezaHistorial.objects.count(), 0)
        self.assert_golden("auditoria", self.client.get("/api/audit/logs/"))

    def test_endpoints_de_materiales_e_historial(self):
        self.assert_golden("materiales", self.client.get(f"{BASE}/pieza-materias-primas/"))
        self.assert_golden("historial", self.client.get(f"{BASE}/pieza-historial/"))


class ImportacionProyectosTests(ApiTestCase):
    URL = f"{BASE}/proyectos/import"

    CSV_MIXTO = (
        "Codigo,Nombre del proyecto,Descripcion,Inicio,Fin,Estado\n"
        "PRJ-010,Parque Norte,Toboganes,2026-06-01,2026-12-15,Planeación\n"
        "aq-01,Código repetido,,,,activo\n"
        "PRJ-010,Repetido en archivo,,,,\n"
        ",Fechas al revés,,2026-06-10,2026-06-01,activo\n"
        ",Sin estado válido,,no-es-fecha,,cerrado\n"
        ",,,,,\n"
        "PRJ-011,,Sin nombre,,,\n"
    )
    CSV_LIMPIO = (
        "codigo,nombre,descripcion,fecha_inicio,fecha_fin,estado\n"
        "PRJ-010,Parque Norte,Toboganes,2026-06-01,2026-12-15,activo\n"
        ",Mantenimiento,,,,\n"
    )

    def test_plantilla(self):
        response = self.client.get(f"{self.URL}/template/")
        filas = [list(fila) for fila in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)]

        self.assert_golden_data("plantilla", {"status": response.status_code, "content_disposition": response["Content-Disposition"], "filas": filas})

    def test_vista_previa_con_errores(self):
        self.assert_golden("vista_previa", self.client.post(f"{self.URL}/preview/", {"file": csv_file(self.CSV_MIXTO, "proyectos.csv")}, format="multipart"))

    def test_confirmar_crea_los_proyectos(self):
        self.assert_golden("confirmar", self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO, "proyectos.csv")}, format="multipart"))
        self.assert_golden("proyectos", self.client.get(f"{BASE}/proyectos/"))

    def test_confirmar_con_errores_no_crea_nada(self):
        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_MIXTO, "proyectos.csv")}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["created_count"], 0)
        self.assertEqual(len(self.client.get(f"{BASE}/proyectos/").json()), 2)

    def test_solo_el_administrador_importa(self):
        self.login_as(self.trabajador)

        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO, "proyectos.csv")}, format="multipart")

        self.assertEqual(response.status_code, 403)


class ImportacionOrdenesTests(ApiTestCase):
    URL = f"{BASE}/ordenes/import"

    CSV_MIXTO = (
        "Proyecto,Orden,Fecha,Estado\n"
        "AQ-01,ORD-010,2026-03-01,En proceso\n"
        "aquaventura,ord-001,,\n"
        "Sin ordenes,ORD-001,,cerrada\n"
        "Sin ordenes,ORD-001,,\n"
        "Kalahari,ORD-020,,abierta\n"
        "AQ-01,,31/02/2026,terminada\n"
        ",,,\n"
    )
    CSV_LIMPIO = "proyecto,codigo_orden,fecha_creacion,estado\nAQ-01,ORD-010,2026-03-01,en_proceso\nSin órdenes,ORD-001,,\n"

    def test_plantilla(self):
        response = self.client.get(f"{self.URL}/template/")
        filas = [list(fila) for fila in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)]

        self.assert_golden_data("plantilla", {"status": response.status_code, "content_disposition": response["Content-Disposition"], "filas": filas})

    def test_vista_previa_con_errores(self):
        self.assert_golden("vista_previa", self.client.post(f"{self.URL}/preview/", {"file": csv_file(self.CSV_MIXTO, "ordenes.csv")}, format="multipart"))

    def test_confirmar_crea_las_ordenes_en_su_proyecto(self):
        self.assert_golden("confirmar", self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO, "ordenes.csv")}, format="multipart"))
        self.assert_golden("ordenes", self.client.get(f"{BASE}/ordenes/"))

    def test_confirmar_con_errores_no_crea_nada(self):
        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_MIXTO, "ordenes.csv")}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["created_count"], 0)
        self.assertEqual(len(self.client.get(f"{BASE}/ordenes/").json()), 2)

    def test_solo_el_administrador_importa(self):
        self.login_as(self.almacenista)

        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO, "ordenes.csv")}, format="multipart")

        self.assertEqual(response.status_code, 403)


class ImportacionPiezasTests(ApiTestCase):
    URL = f"{BASE}/piezas/import"

    CSV_MIXTO = (
        "Trace,Nombre pieza,Proyecto,Orden,Responsable,Gelcoat,QC,Peso,Materias primas\n"
        "TR-0100,Tobogán C,AQ-01,ORD-001,carlos@innolution.test,2026-06-05,,48.2,Resina poliester: 20; Fibra de vidrio = 10\n"
        "tr-0001,Repetida,Aquaventura,ORD-001,Juan Administrador,2026-06-05,,,\n"
        "TR-0100,Repetida en archivo,AQ-01,ORD-002,Carla Mejía,2026-06-05,,,\n"
        "TR-0101,Orden ajena,Sin ordenes,ORD-001,carlos@innolution.test,2026-06-05,,,\n"
        "TR-0102,Usuario y proyecto,Atlantis,ORD-001,nadie@innolution.test,,no-fecha,-3,\n"
        "TR-0103,Materiales,AQ-01,ORD-001,Carlos,2026-06-05,,,Catalizador MEKP: 1; Resina poliéster: 0; Resina poliéster: 2; Kevlar: 3; Gelcoat blanco\n"
        ",,,,,,,,\n"
    )
    CSV_LIMPIO = (
        "trace_id,nombre,proyecto,orden,usuario,fecha_gelcoat,fecha_qc,peso_real,materiales\n"
        "TR-0100,Tobogán C,AQ-01,ORD-002,carlos@innolution.test,2026-06-05,2026-06-12,48.2,Resina poliéster: 20; Fibra de vidrio: 10\n"
        "TR-0101,Tobogán D,Aquaventura,ORD-001,Juan Administrador,2026-06-06,,,\n"
    )

    def test_plantilla(self):
        response = self.client.get(f"{self.URL}/template/")
        filas = [list(fila) for fila in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)]

        self.assert_golden_data("plantilla", {"status": response.status_code, "content_disposition": response["Content-Disposition"], "filas": filas})

    def test_vista_previa_con_errores(self):
        self.assert_golden("vista_previa", self.client.post(f"{self.URL}/preview/", {"file": csv_file(self.CSV_MIXTO, "piezas.csv")}, format="multipart"))

    def test_el_trabajador_importa_y_queda_en_el_historial(self):
        self.login_as(self.trabajador)

        self.assert_golden("confirmar", self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO, "piezas.csv")}, format="multipart"))
        self.assert_golden("piezas", self.client.get(f"{BASE}/piezas/"))

        historial = PiezaHistorial.objects.get(pieza__trace_id="TR-0100")
        self.assertEqual((historial.accion, historial.usuario_id), ("creacion", self.trabajador.id))

    def test_confirmar_con_errores_no_crea_nada(self):
        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_MIXTO, "piezas.csv")}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["created_count"], 0)
        self.assertEqual(Pieza.objects.count(), 1)
        self.assertEqual(PiezaMateriaPrima.objects.count(), 2)

    def test_el_almacenista_no_puede_importar(self):
        self.login_as(self.almacenista)

        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO, "piezas.csv")}, format="multipart")

        self.assertEqual(response.status_code, 403)
