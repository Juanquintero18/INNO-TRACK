"""Pruebas de caracterización de la API de inventario."""

from datetime import date
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.inventory.models import MateriaPrimaPiezaConfig, MovimientoInventario
from config.testing import ApiTestCase


BASE = "/api/inventory"


class CatalogosTests(ApiTestCase):
    def test_listados(self):
        for recurso in ("unidades-medida", "proveedores", "trabajadores", "materias-primas", "movimientos"):
            self.assert_golden(f"listado_{recurso}", self.client.get(f"{BASE}/{recurso}/"))

    def test_requiere_sesion(self):
        self.login_as(None)

        self.assert_golden("sin_sesion", self.client.get(f"{BASE}/proveedores/"))

    def test_permisos_de_escritura_por_rol(self):
        intentos = {
            "proveedores": {"nombre": "Nuevo proveedor"},
            "trabajadores": {"codigo_trabajador": "TRB-099", "nombre": "Nuevo"},
            "unidades-medida": {"nombre": "Galón", "abreviatura": "gal"},
            "materias-primas": {"nombre": "Pigmento"},
        }
        resultado = {}

        for rol, usuario in (("almacenista", self.almacenista), ("trabajador", self.trabajador)):
            self.login_as(usuario)
            for recurso, payload in intentos.items():
                response = self.client.post(f"{BASE}/{recurso}/", payload, format="json")
                resultado[f"{rol} crea {recurso}"] = response.status_code

        self.assert_golden_data("matriz", resultado)

    def test_crud_de_proveedor_con_auditoria(self):
        creado = self.client.post(f"{BASE}/proveedores/", {"nombre": "Pinturas Norte", "telefono": "6045551212", "email": "hola@pinturas.test", "direccion": "Medellín"}, format="json")
        self.assert_golden("crear", creado)

        url = f"{BASE}/proveedores/{creado.json()['id']}/"
        self.assert_golden("actualizar", self.client.put(url, {"nombre": "Pinturas del Norte"}, format="json"))
        self.assert_golden("eliminar", self.client.delete(url))
        self.assert_golden("auditoria_tras_eliminar", self.client.get("/api/audit/logs/"))

    def test_crud_de_trabajador(self):
        creado = self.client.post(f"{BASE}/trabajadores/", {"codigo_trabajador": "TRB-020", "nombre": "Nora Paz"}, format="json")
        self.assert_golden("crear", creado)
        self.assert_golden("eliminar", self.client.delete(f"{BASE}/trabajadores/{creado.json()['id']}/"))


class MateriasPrimasTests(ApiTestCase):
    def test_crear_actualizar_y_eliminar(self):
        creada = self.client.post(f"{BASE}/materias-primas/", {"nombre": "Resina viniléster", "unidad_medida_id": 2, "costo": 18.4, "fecha_actualizacion": "2026-03-14"}, format="json")
        self.assert_golden("crear", creada)

        url = f"{BASE}/materias-primas/{creada.json()['id']}/"
        self.assert_golden("actualizar", self.client.put(url, {"nombre": "Resina viniléster premium", "unidad_medida_id": 1, "costo": 20, "fecha_actualizacion": "2026-03-15"}, format="json"))
        self.assert_golden("eliminar", self.client.delete(url))

    def test_umbrales_de_estabilidad(self):
        url = f"{BASE}/materias-primas/{self.resina.id}/stability-thresholds/"

        self.assert_golden("guardar", self.client.put(url, {"stock_critico_max": 15, "stock_bajo_max": 40.5}, format="json"))
        self.assert_golden("bajo_menor_que_critico", self.client.put(url, {"stock_critico_max": 50, "stock_bajo_max": 40}, format="json"))
        self.assert_golden("critico_negativo", self.client.put(url, {"stock_critico_max": -1, "stock_bajo_max": 40}, format="json"))

        self.login_as(self.trabajador)
        self.assert_golden("sin_permiso", self.client.put(url, {"stock_critico_max": 1, "stock_bajo_max": 2}, format="json"))

    def test_materiales_habilitados_para_piezas(self):
        url = f"{BASE}/materias-primas/piezas-materiales/"

        self.assert_golden("habilitar_dos", self.client.put(url, {"enabled_materia_ids": [self.resina.id, self.catalizador.id]}, format="json"))
        self.assertEqual(
            sorted(MateriaPrimaPiezaConfig.objects.values_list("materia_prima_id", "enabled_for_piezas")),
            [(self.fibra.id, False), (self.gelcoat.id, False)],
        )

        self.assert_golden("ids_inexistentes", self.client.put(url, {"enabled_materia_ids": [self.resina.id, 998, 999]}, format="json"))
        self.assert_golden("lista_vacia", self.client.put(url, {"enabled_materia_ids": []}, format="json"))

        self.login_as(self.almacenista)
        self.assert_golden("solo_admin", self.client.put(url, {"enabled_materia_ids": [self.resina.id]}, format="json"))


class MovimientosTests(ApiTestCase):
    def test_crear_asigna_el_usuario_de_la_sesion(self):
        self.login_as(self.almacenista)

        response = self.client.post(f"{BASE}/movimientos/", {"materia_prima_id": self.gelcoat.id, "proveedor_id": self.fibramax.id, "trabajador_produccion_id": None, "tipo": "entrada", "cantidad": 12.5, "fecha": "2026-03-14", "motivo": None, "referencia": "OC-2001"}, format="json")

        self.assert_golden("crear", response)

    def test_no_permite_dejar_el_stock_en_negativo(self):
        salida = {"materia_prima_id": self.gelcoat.id, "tipo": "salida", "cantidad": 9, "fecha": "2026-03-14", "referencia": "PRD-0100"}
        self.assert_golden("salida_excesiva", self.client.post(f"{BASE}/movimientos/", salida, format="json"))

        ajuste = {"materia_prima_id": self.gelcoat.id, "tipo": "ajuste", "cantidad": -8.01, "fecha": "2026-03-14", "referencia": "AJU-1"}
        self.assert_golden("ajuste_excesivo", self.client.post(f"{BASE}/movimientos/", ajuste, format="json"))

        self.assert_golden("salida_valida", self.client.post(f"{BASE}/movimientos/", {**salida, "cantidad": 8}, format="json"))

    def test_actualizar_recalcula_el_stock_sin_contar_el_propio_movimiento(self):
        salida = MovimientoInventario.objects.get(referencia="PRD-0007")
        url = f"{BASE}/movimientos/{salida.id}/"
        payload = {"materia_prima_id": self.resina.id, "tipo": "salida", "fecha": "2026-01-12", "referencia": "PRD-0007", "trabajador_produccion_id": self.pedro.id}

        self.assert_golden("dentro_del_stock", self.client.put(url, {**payload, "cantidad": 95}, format="json"))
        self.assert_golden("fuera_del_stock", self.client.put(url, {**payload, "cantidad": 95.01}, format="json"))

    def test_eliminar_deja_auditoria(self):
        movimiento = MovimientoInventario.objects.get(referencia="OC-1002")

        self.assert_golden("eliminar", self.client.delete(f"{BASE}/movimientos/{movimiento.id}/"))
        self.assert_golden("auditoria", self.client.get("/api/audit/logs/"))


def csv_file(text, name="movimientos.csv"):
    return SimpleUploadedFile(name, text.encode("utf-8"), content_type="text/csv")


class ImportacionTests(ApiTestCase):
    # Encabezados con alias, formatos de número y fecha variados, y varias filas con errores.
    CSV_MIXTO = (
        "Materia Prima,Tipo movimiento,Cantidad,Fecha,Proveedor,Responsable,Motivo,Codigo referencia\n"
        "resina poliester,Entrada,\"25,5\",12/03/2026,FibraMax,,Compra,OC-3001\n"
        "Fibra de vidro,salida,10,2026-03-12,,trb-001,Consumo,PRD-3002\n"
        "Gelcoat blanco,salida,100,2026-03-12,,Pedro Gomez,,PRD-3003\n"
        "Gelcoat blanco,traslado,abc,31-31-2026,Proveedor X,Nadie,,\n"
        "\n"
        "Catalizador MEKP,ajuste,0,2026/03/12,,,Conteo,AJU-3004\n"
    )

    CSV_LIMPIO = (
        "materia_prima,tipo,cantidad,fecha,proveedor,trabajador,motivo,referencia\n"
        "Resina poliéster,entrada,25.5,2026-03-12,FibraMax,,Compra,OC-4001\n"
        "Resina poliéster,salida,10,2026-03-13,,TRB-002,,PRD-4002\n"
        "Fibra de vidrio,ajuste,-2,2026-03-13,,,Conteo,AJU-4003\n"
    )

    def test_plantilla(self):
        response = self.client.get(f"{BASE}/movimientos/import/template/")
        hoja = load_workbook(BytesIO(response.content)).active

        self.assert_golden_data(
            "plantilla",
            {
                "status": response.status_code,
                "content_type": response["Content-Type"],
                "content_disposition": response["Content-Disposition"],
                "hoja": hoja.title,
                "filas": [list(fila) for fila in hoja.iter_rows(values_only=True)],
            },
        )

    def test_vista_previa_con_encabezados_alternativos_y_errores(self):
        response = self.client.post(f"{BASE}/movimientos/import/preview/", {"file": csv_file(self.CSV_MIXTO)}, format="multipart")

        self.assert_golden("vista_previa", response)
        self.assertEqual(MovimientoInventario.objects.count(), 5)

    def test_vista_previa_de_excel(self):
        libro = Workbook()
        hoja = libro.active
        hoja.append(["materia-prima", "tipo", "cantidad", "fecha", "nombre proveedor", "trabajador_codigo", "motivo", "referencia"])
        hoja.append(["Gelcoat blanco", "entrada", 4.25, date(2026, 3, 12), "Química Industrial S.A.", None, None, "OC-5001"])
        contenido = BytesIO()
        libro.save(contenido)
        archivo = SimpleUploadedFile("movimientos.xlsx", contenido.getvalue())

        self.assert_golden("vista_previa_excel", self.client.post(f"{BASE}/movimientos/import/preview/", {"file": archivo}, format="multipart"))

    def test_archivos_invalidos(self):
        url = f"{BASE}/movimientos/import/preview/"

        self.assert_golden("sin_archivo", self.client.post(url, {}, format="multipart"))
        self.assert_golden("extension_no_soportada", self.client.post(url, {"file": csv_file("a,b", name="datos.txt")}, format="multipart"))
        self.assert_golden("solo_encabezados", self.client.post(url, {"file": csv_file("materia_prima,tipo,cantidad,fecha,referencia\n")}, format="multipart"))
        self.assert_golden("vacio", self.client.post(url, {"file": csv_file("")}, format="multipart"))

    def test_confirmar_importacion_crea_los_movimientos(self):
        self.login_as(self.almacenista)

        response = self.client.post(f"{BASE}/movimientos/import/commit/", {"file": csv_file(self.CSV_LIMPIO)}, format="multipart")

        self.assert_golden("confirmar", response)
        self.assert_golden("movimientos_creados", self.client.get(f"{BASE}/movimientos/"))

    def test_confirmar_con_errores_no_crea_nada(self):
        response = self.client.post(f"{BASE}/movimientos/import/commit/", {"file": csv_file(self.CSV_MIXTO)}, format="multipart")

        self.assert_golden("rechazada", response)
        self.assertEqual(MovimientoInventario.objects.count(), 5)

    def test_solo_admin_y_almacenista_importan(self):
        self.login_as(self.trabajador)

        response = self.client.post(f"{BASE}/movimientos/import/commit/", {"file": csv_file(self.CSV_LIMPIO)}, format="multipart")

        self.assertEqual(response.status_code, 403)
        self.assertEqual(MovimientoInventario.objects.count(), 5)


class ImportacionProveedoresTests(ApiTestCase):
    URL = f"{BASE}/proveedores/import"

    CSV_MIXTO = (
        "Razon social,Celular,Correo electronico,Domicilio\n"
        "Pinturas Norte,6045551212,hola@pinturasnorte.test,Medellín\n"
        "Duplicado existente,3001112233,VENTAS@quimica.test,Cali\n"
        "Repetido en archivo,3001112234,hola@pinturasnorte.test,Cali\n"
        ",,no-es-correo,\n"
    )

    CSV_LIMPIO = (
        "nombre,telefono,email,direccion\n"
        "Pinturas Norte,6045551212,hola@pinturasnorte.test,Medellín\n"
        "Lijas y Abrasivos,3105550000,ventas@lijas.test,Bogotá\n"
    )

    def test_plantilla(self):
        response = self.client.get(f"{self.URL}/template/")
        hoja = load_workbook(BytesIO(response.content)).active

        self.assert_golden_data(
            "plantilla",
            {"status": response.status_code, "content_disposition": response["Content-Disposition"], "filas": [list(fila) for fila in hoja.iter_rows(values_only=True)]},
        )

    def test_vista_previa_con_errores(self):
        self.assert_golden("vista_previa", self.client.post(f"{self.URL}/preview/", {"file": csv_file(self.CSV_MIXTO)}, format="multipart"))

    def test_confirmar_crea_los_proveedores(self):
        self.login_as(self.almacenista)

        self.assert_golden("confirmar", self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO)}, format="multipart"))
        self.assert_golden("proveedores", self.client.get(f"{BASE}/proveedores/"))

    def test_confirmar_con_errores_no_crea_nada(self):
        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_MIXTO)}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["created_count"], 0)
        self.assertEqual(len(self.client.get(f"{BASE}/proveedores/").json()), 2)

    def test_el_trabajador_no_puede_importar(self):
        self.login_as(self.trabajador)

        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO)}, format="multipart")

        self.assertEqual(response.status_code, 403)


class ImportacionTrabajadoresTests(ApiTestCase):
    URL = f"{BASE}/trabajadores/import"

    CSV_MIXTO = (
        "Codigo,Nombre completo\n"
        "TRB-020,Nora Paz\n"
        "trb-001,Ya existe\n"
        "TRB-020,Repetido\n"
        ",\n"
        "TRB-021,\n"
    )

    def test_plantilla(self):
        response = self.client.get(f"{self.URL}/template/")
        filas = [list(fila) for fila in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)]

        self.assert_golden_data("plantilla", {"status": response.status_code, "content_disposition": response["Content-Disposition"], "filas": filas})

    def test_vista_previa_con_errores(self):
        self.assert_golden("vista_previa", self.client.post(f"{self.URL}/preview/", {"file": csv_file(self.CSV_MIXTO)}, format="multipart"))

    def test_confirmar_crea_los_trabajadores(self):
        limpio = "codigo_trabajador,nombre\nTRB-020,Nora Paz\nTRB-021,Iván Mora\n"

        self.assert_golden("confirmar", self.client.post(f"{self.URL}/commit/", {"file": csv_file(limpio)}, format="multipart"))
        self.assert_golden("trabajadores", self.client.get(f"{BASE}/trabajadores/"))

    def test_solo_el_administrador_importa(self):
        self.login_as(self.almacenista)

        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file("codigo_trabajador,nombre\nTRB-020,Nora\n")}, format="multipart")

        self.assertEqual(response.status_code, 403)


class ImportacionMateriasPrimasTests(ApiTestCase):
    URL = f"{BASE}/materias-primas/import"

    CSV_MIXTO = (
        "Material,Unidad,Costo unitario,Fecha,Critico hasta,Bajo hasta\n"
        "Gelcoat gris,L,21.00,2026-05-02,,\n"
        "resina poliéster,kg,10,2026-05-02,,\n"
        "Gelcoat gris,litro,20,2026-05-02,,\n"
        "Pigmento,toneladas,0,,,\n"
        "Malla,kilogramo,5,2026-05-02,30,10\n"
        "Cera,kg,4,02/05/2026,5,\n"
        ",,,,,\n"
    )
    CSV_LIMPIO = (
        "nombre,unidad_medida,costo,fecha_actualizacion,stock_critico_max,stock_bajo_max\n"
        "Gelcoat gris,Litro,21.00,2026-05-02,,\n"
        "Resina viniléster,kg,18.40,2026-05-02,20,50\n"
    )

    def test_plantilla(self):
        response = self.client.get(f"{self.URL}/template/")
        filas = [list(fila) for fila in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)]

        self.assert_golden_data("plantilla", {"status": response.status_code, "content_disposition": response["Content-Disposition"], "filas": filas})

    def test_vista_previa_con_errores(self):
        self.assert_golden("vista_previa", self.client.post(f"{self.URL}/preview/", {"file": csv_file(self.CSV_MIXTO)}, format="multipart"))

    def test_confirmar_crea_las_materias_con_sus_umbrales(self):
        self.login_as(self.almacenista)

        self.assert_golden("confirmar", self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO)}, format="multipart"))
        self.assert_golden("materias", self.client.get(f"{BASE}/materias-primas/"))

    def test_confirmar_con_errores_no_crea_nada(self):
        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_MIXTO)}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["created_count"], 0)
        self.assertEqual(len(self.client.get(f"{BASE}/materias-primas/").json()), 4)

    def test_el_trabajador_no_puede_importar(self):
        self.login_as(self.trabajador)

        response = self.client.post(f"{self.URL}/commit/", {"file": csv_file(self.CSV_LIMPIO)}, format="multipart")

        self.assertEqual(response.status_code, 403)
